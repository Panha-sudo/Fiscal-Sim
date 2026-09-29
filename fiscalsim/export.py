"""Excel workbook of scenario results, for officials who work in spreadsheets."""
from __future__ import annotations

import io
from dataclasses import asdict

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as C
from .report import LABELS

HEADER = PatternFill("solid", fgColor="E1EFEB")


def _about(runs: int, source: str) -> pd.DataFrame:
    rows = [
        ("Model", "AI fiscal simulation of Cambodia's civil service pay and pensions (thesis prototype)"),
        ("Data", source),
        ("Base year / horizon", f"{C.BASE_YEAR} to {C.END_YEAR}"),
        ("Monte Carlo runs", f"{runs:,}"),
        ("Bands", "p5 and p95 are the 5th and 95th percentiles across runs; p50 is the median"),
        ("Total cost", "Wage bill + government pension cost (employer contribution + top-up once the fund is empty)"),
        ("Replacement rate", "First pension / final total pay (basic + allowances), central path"),
        ("Ministry and province projections", "Each unit keeps its base-year share of its sector's wage bill"),
    ]
    return pd.DataFrame(rows, columns=["Item", "Value"])


def scenario_levers(scenarios: dict) -> pd.DataFrame:
    return pd.DataFrame([asdict(sc) for sc in scenarios.values()]).set_index("code")


def _fit(ws):
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


def workbook(frames: dict, scenarios: dict, runs: int, source: str,
             ministry: pd.DataFrame | None = None, province: pd.DataFrame | None = None,
             extra: dict | None = None) -> bytes:
    """frames: output of simulate.results_frames. Returns .xlsx bytes."""
    fans = frames["fans"].copy()
    fans["indicator"] = fans["metric"].map(LABELS)
    sheets = {
        "About": _about(runs, source),
        "Summary": frames["table"].reset_index(),
        "Scenario levers": scenario_levers(scenarios).reset_index(),
        "Yearly bands": fans[["scenario", "indicator", "year", "p5", "p50", "p95", "central"]],
        "Adequacy": frames["adequacy"].reset_index(),
        "Workforce": frames["workforce"].reset_index(),
        "Wage bill by sector": (frames["wage_bill_sector"] / 1e9).round(1).reset_index()
                               .rename(columns=lambda c: c if c in ("scenario", "year") else f"{c} (bn riel)"),
    }
    if ministry is not None:
        sheets["By ministry"] = ministry.reset_index()
    if province is not None:
        sheets["By province"] = province.reset_index()
    for name, df in (extra or {}).items():
        sheets[name[:31]] = df.reset_index() if df.index.name else df
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            _fit(xw.sheets[name])
    return buf.getvalue()
